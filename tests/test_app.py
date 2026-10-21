"""Web app, file readers and report builders. Run: python -m pytest -q"""
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.main import MAX_ROWS, app
from phone_validator.readers import read_upload
from phone_validator.report import build_docx, build_pdf
from phone_validator.validate import validate_table

ROOT = Path(__file__).resolve().parent.parent
SAMPLE = ROOT / "data" / "samples" / "synthetic_200.csv"
client = TestClient(app)


# ---------------------------------------------------------------- single number
def test_index_page_served():
    r = client.get("/")
    assert r.status_code == 200 and "Champ Phone Validator" in r.text


@pytest.mark.parametrize("number", ["+91 98450 12345", "09845012345", "98450-12345", "919845012345"])
def test_check_valid_indian_mobile_is_validated(number):
    r = client.post("/api/check", json={"number": number, "country": "IN"})
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "Validated" and body["validated"] is True
    assert body["e164"] == "+919845012345" and body["international"] == "+91 98450 12345"
    assert body["line_type"] == "mobile" and body["carrier"]


@pytest.mark.parametrize("number,reason", [
    ("12345", "invalid_too_short"),
    ("080 2345 6789", "landline"),
    ("hello", "unparseable"),
    ("+91 5845012345", "invalid_unallocated_range"),
])
def test_check_bad_numbers_are_not_validated(number, reason):
    body = client.post("/api/check", json={"number": number, "country": "IN"}).json()
    assert body["status"] == "Not Validated" and body["validated"] is False
    assert body["reason"] == reason and body["reason_text"]


def test_check_number_with_own_country_code_ignores_picked_country():
    body = client.post("/api/check", json={"number": "+44 7400 123456", "country": "IN"}).json()
    assert body["status"] == "Validated" and body["country"] == "GB"


def test_check_us_number_is_not_validated_but_flagged_for_review():
    body = client.post("/api/check", json={"number": "212 555 0123", "country": "US"}).json()
    assert body["status"] == "Not Validated" and body["needs_review"] is True


def test_check_rejects_empty_and_overlong_input():
    assert client.post("/api/check", json={"number": "  ", "country": "IN"}).status_code == 422
    r = client.post("/api/check", json={"number": "9" * 200, "country": "IN"})
    assert r.status_code == 422 and "error" in r.json()


# ---------------------------------------------------------------- bulk upload
@pytest.mark.parametrize("fmt,magic,ctype", [
    ("pdf", b"%PDF", "application/pdf"),
    ("docx", b"PK", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ("csv", "﻿row,".encode("utf-8"), "text/csv"),
])
def test_bulk_returns_report_in_each_format(fmt, magic, ctype):
    with SAMPLE.open("rb") as fh:
        r = client.post("/api/bulk", files={"file": ("synthetic_200.csv", fh, "text/csv")},
                        data={"format": fmt, "country": "IN"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(ctype)
    assert r.content.startswith(magic)
    assert f'synthetic_200_validation_report.{fmt}' in r.headers["content-disposition"]
    s = json.loads(r.headers["x-validation-summary"])
    assert s == {"total": 200, "validated": 108, "not_validated": 92, "needs_review": 27, "duplicates": 20,
                 "phone_column": "phone", "country_column": "country", "default_country": "IN"}


def test_bulk_xlsx_upload_with_numeric_cells():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Mobile Number"])
    ws.append(["A", 9845012345])           # int cell
    ws.append(["B", 919886011223.0])       # float cell, the classic Excel case
    ws.append(["C", "080 2345 6789"])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/bulk", files={"file": ("list.xlsx", buf.getvalue(), "application/octet-stream")},
                    data={"format": "csv"})
    assert r.status_code == 200, r.text
    s = json.loads(r.headers["x-validation-summary"])
    assert (s["total"], s["validated"], s["not_validated"], s["phone_column"]) == (3, 2, 1, "Mobile Number")


def test_bulk_txt_and_headerless_csv():
    txt = "phone\n+91 98450 12345\n09886011223\nnot a number\n"
    r = client.post("/api/bulk", files={"file": ("n.txt", txt.encode(), "text/plain")}, data={"format": "csv"})
    s = json.loads(r.headers["x-validation-summary"])
    assert (s["total"], s["validated"]) == (3, 2)
    csv_no_header = "9845012345\n9886011223\n12345\n"
    r = client.post("/api/bulk", files={"file": ("n.csv", csv_no_header.encode(), "text/csv")}, data={"format": "csv"})
    s = json.loads(r.headers["x-validation-summary"])
    assert (s["total"], s["validated"]) == (3, 2)


@pytest.mark.parametrize("name,content,status,needle", [
    ("x.pdf", b"%PDF-1.4", 415, "Unsupported file type"),
    ("x.csv", b"", 422, "empty"),
    ("x.csv", b"name,email\nA,a@b.c\n", 422, "phone column"),
    ("x.xlsx", b"not really excel", 422, "Could not read"),
])
def test_bulk_errors_are_readable(name, content, status, needle):
    r = client.post("/api/bulk", files={"file": (name, content, "application/octet-stream")}, data={"format": "pdf"})
    assert r.status_code == status and needle in r.json()["error"]


def test_bulk_rejects_bad_format_and_too_many_rows():
    r = client.post("/api/bulk", files={"file": ("x.csv", b"phone\n9845012345\n", "text/csv")}, data={"format": "xls"})
    assert r.status_code == 422
    big = ("phone\n" + "9845012345\n" * (MAX_ROWS + 1)).encode()
    r = client.post("/api/bulk", files={"file": ("big.csv", big, "text/csv")}, data={"format": "csv"})
    assert r.status_code == 413 and "limit" in r.json()["error"]


def test_sample_download():
    r = client.get("/api/sample")
    assert r.status_code == 200 and r.text.startswith("id,name,phone,country")


# ---------------------------------------------------------------- readers and reports
def test_semicolon_csv_is_sniffed():
    rows, fields = read_upload("x.csv", "Name;Phone;Country\nA;+91 98450 12345;India\n".encode())
    assert fields == ["Name", "Phone", "Country"] and rows[0]["Phone"] == "+91 98450 12345"


def test_reports_contain_summary_and_survive_odd_characters():
    rows = [{"phone": p} for p in ["+91 98450 12345", "‐98450–12345’", "\x0bbad\x00", "080 2345 6789"]]
    table = validate_table(rows, ["phone"])
    pdf = build_pdf(table, "odd.csv")
    docx = build_docx(table, "odd.csv")
    assert pdf.startswith(b"%PDF") and len(pdf) > 2000
    from docx import Document
    text = "\n".join(p.text for p in Document(io.BytesIO(docx)).paragraphs)
    assert "Phone Number Validation Report" in text and "odd.csv" in text
